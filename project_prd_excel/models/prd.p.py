from datetime import datetime, timedelta 
from odoo import api, fields, models, _
from odoo.exceptions import UserError, ValidationError, AccessError
import logging


import openpyxl
import base64
from io import BytesIO



_logger = logging.getLogger(__name__)


class ProductRequirementDocument(models.Model):
    _inherit = 'prd.document'
    
    def action_excel_wizard(self):
        return {
            'type': 'ir.actions.act_window',
            'name': 'Load Requirements from Excel',
            'res_model': 'prd.excel.wizard',
            'view_mode': 'form',
            'target': 'new',
            # ~ 'context': {'default_week_template_id': self.env.context["default_week_template_id"]},
        }
        return action



class ExcelWizard(models.TransientModel):
    _name = 'prd.excel.wizard'
    _description = 'Load Requirements from Excel'

    file = fields.Binary(string="File", required=True)
 
 

    def import_excel(self):

            wb = openpyxl.load_workbook(filename=BytesIO(base64.b64decode(self.file)))
            
            # ~ raise UserError(f"{wb.sheetnames=}")
            for ws in wb.sheetnames:
                activews = wb[ws]
                record = {}                
                for row in range(1, activews.max_row + 1):
                    value = activews.cell(row=row,column=1).value
                    if value and any(c.isdigit() for c in str(value)): # Has numbers
                        record = {
                            'page': ws,
                            'no': activews.cell(row=row,column=1),
                             'category':  activews.cell(row=row,column=2).value if len(activews.cell(row=row,column=2).value) < 25 else '',
                             'name': activews.cell(row=row,column=2) if len(activews.cell(row=row,column=2).value) > 25 else activews.cell(row=row,column=3).value ,
                            'prd_id': self.env.context['active_id'],
                        
                        }
                        self.env['prd.requirement'].create(record)

                    # ~ if activews.cell(row=row,column=1) != None:
                        # ~ for col in range(1, 13):
                            # ~ record_value = activews.cell(row=row, column=col).value
                            # ~ if record_value != None:
                                # ~ record_value = str(record_value).strip()
                                # ~ record_value = record_value.lower()
                                # ~ if record_value == "monerary":
                                    # ~ record_value = "monetary"
                                # ~ if "percentage" in record_value:
                                    # ~ record_value = record_value.replace("percentage", "percent")
                            # ~ record[self.field_keys[col-1]] = record_value
                        # ~ record['csrd_sheet_name'] = ws
                        # ~ self.create_record(record)
            # ~ return {
                # ~ 'type': 'ir.actions.client',
                # ~ 'tag': 'reload',
            # ~ }


    def create_record(self,record):

            csrd_esrs_id = self.env["csrd.esrs"].search([("csrd_id", "=", record["csrd_id"])])

            if not csrd_esrs_id:

                self.env["csrd.esrs"].create(record)



