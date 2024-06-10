from odoo import api, fields, models, _
from odoo.exceptions import ValidationError, UserError
import logging
import requests
import re
import os

from openpyxl import load_workbook
from xlrd import open_workbook, XLRDError
from xlrd.book import Book
from xlrd.sheet import Sheet
import base64
from io import BytesIO

_logger = logging.getLogger(__name__)


class AddStakeholder(models.TransientModel):
    _name = "project.add.stakeholder.wizard"
    _description = "Add stakeholders to Projects."

    partner_id = fields.Many2one(comodel_name='res.partner', string="Partner",
                                 help="Product owner / customer")  # example SKF, Skogsstyrelsen, Dollarstore
    stakeholder_file = fields.Binary(string='Stakeholder modules', help='Excel file exported from an Odoo instance.')
    exclude_odoosa = fields.Boolean(string='Exclude Odoo SA', help='Dont try to check Odoo Core modules.')
    odoo_server = fields.Char(string='Odoo Server',
                              help='Odoo server to check modules/repos, server need to have ssh-key fpr the odoo server')
    message_box = fields.Text(string='')
    project_id = fields.Many2one(comodel_name="project.project", default=lambda b: b.env.context.get('active_id'))
    addrepos_button = fields.Boolean()

    def load_file(self):
        module_file = load_workbook(filename=BytesIO(base64.b64decode(self.stakeholder_file))).active
        author_pos = module_pos = None
        for i in range(1, 20):
            if module_file.cell(1, i).value in ['Författare', 'Author']:
                author_pos = i
            if module_file.cell(1, i).value in ['Tekniskt namn', 'Technical Name']:
                module_pos = i
        if not module_pos:
            raise UserError("Missing Technical Name in file")

        failed_modules = []
        for row in range(2, len(tuple(module_file.rows))):
            author_name = module_file.cell(row, author_pos).value
            module_name = module_file.cell(row, module_pos).value
            if self.exclude_odoosa and author_name in ["Odoo S.A.", "Odoo SA"]:
                continue
            # if we have the module (project.task) add stakeholder
            module = self.env['project.task'].search(
                [('project_id', '=', self.project_id.id), ('git_module', '=', module_name)], limit=1)  # git_module
            print(module_name, module)
            if not module:
                failed_modules.append((module_name, author_name))
            else:
                module.module_stakeholder_ids = [(4, self.partner_id.id, 0)]

        if len(failed_modules) > 0:
            self.message_box = "Failed modules: " + ','.join([str(t) for t in failed_modules])
        return {
            "type": "ir.actions.act_window",
            "name": "Add Stakeholder",
            "res_model": "project.add.stakeholder.wizard",
            "res_id": self.id,
            "view_mode": "form",
            "context": dict(self.env.context),
        }
        # gi

    def _extract_paths(self, file_path):
        # Regex pattern to match the first two directories before __manifest__.py
        pattern = r'([^/]+/[^/]+)/__manifest__.py'

        # Search for the pattern in the given file path
        match = re.search(pattern, file_path)
        result = match.group(1)
        repo, _module = result.split('/')
        file_to_repo = file_path.split(f"/{_module}")[0]
        if match:
            return file_to_repo, repo, _module
        else:
            return None, None

    def check_repos(self):
        module_file = load_workbook(filename=BytesIO(base64.b64decode(self.stakeholder_file))).active
        author_pos = module_pos = 1
        for i in range(1, 20):
            if module_file.cell(1, i).value in ['Författare', 'Author']:
                author_pos = i
            if module_file.cell(1, i).value in ['Tekniskt namn', 'Technical Name']:
                module_pos = i

        failed_modules = []
        for row in range(2, len(tuple(module_file.rows))):
            author_name = module_file.cell(row, author_pos).value
            module_name = module_file.cell(row, module_pos).value
            if self.exclude_odoosa and author_name in ["Odoo S.A.", "Odoo SA"]:
                continue
            # if we have the module (project.task) add stakeholder
            module = self.env['project.task'].search([
                ('project_id', '=', self.project_id.id), ('git_module', '=', module_name)
            ], limit=1)  # git_module
            if not module:
                failed_modules.append((module_name, author_name))
            else:
                module.module_stakeholder_ids = [(4, self.partner_id.id, 0)]

        if len(failed_modules) > 0:
            repos = set()
            repos_dict = {}
            for module in failed_modules:
                base_path = "/home/ayomir/odoo/14.0"

                # if module[1] in ['Odoo S.A.', 'Odoo']:
                #     continue


                command = f'find {base_path} -type f -wholename "*/{module[0]}/__manifest__.py"'
                #repo = os.popen(f'locate {module[0]}/__manifest__.py').read().split('/')
                module_path = os.popen(command).read().splitlines()

                if not module_path:
                    continue

                file_to_repo, repo, _module = self._extract_paths(module_path[0])

                if len(repo) < 3:
                    _logger.info(f"{repo=}")
                    continue
                if repo == 'odootools':
                    continue
                try:
                    if repo not in repos_dict:
                        repos_dict[repo] = [self.env['project.project'].get_git_origin(file_to_repo), [_module]]
                    else:
                        repos_dict[repo][1].append(_module)
                    repos.add(f"{[repo]}{_module}\n")
                except Exception as e:
                    raise UserError(repo)

            # self.message_box = "Missing repos: " + ','.join(sorted(repos))
            message_box = ""
            for keys, vals in repos_dict.items():
                message_box += f"{keys}:origin {vals[0]},modules: {vals[1]}\n"
            self.message_box = message_box
        self.addrepos_button = True
        return {
            "type": "ir.actions.act_window",
            "name": "Add Stakeholder",
            "res_model": "project.add.stakeholder.wizard",
            "res_id": self.id,
            "view_mode": "form",
            # ~ "domain": [("journal_id", "=", self.id)],
            "context": dict(self.env.context),
        }

        # git_module i stället för name

        # if we don't have the module add to list of missing modules
        # What will we do with Odoo Core?

    def add_repos(self):
        missing_modules = []
        # TODO Add Repos

        message_box = list(filter(None, self.message_box.split('\n')))
        for module_n_repo in message_box:
            repo_n_repo_url, modules = module_n_repo.split(',modules: ')
            repo_name, repo_url = repo_n_repo_url.split(':origin ')

            if self.env['project.project'].is_valid_git_url(repo_url):
                repo_author = repo_url.split('git@github.com:')[-1].split('/')[0]
                print("repo_author", repo_author, "repo_name", repo_name)

                # for module in
                self.env['git.repos'].load_modules(
                    author=repo_author, repo_name=repo_name, project_id=self.project_id.id
                )

        # failed_modules = self.message_box.split('Failed modules: ')[1]
        # components = re.findall(r"\('.*?'\)", failed_modules)
        # module_tuple = [eval(component) for component in components]

            # for module_technical_name, module_author in module_tuple:
            #     res = self.env['git.repos'].load_modules(module_author, module_technical_name, self.project_id.id, False,
            #                                              stakeholder=self.partner_id)
            #     print("res", res)
            # for m in res:
            #     missing_modules.append(m)

        # self.message_box = f"Missing modules {missing_modules}"
        self.addrepos_button = False

        return {
            "type": "ir.actions.act_window",
            "name": "Add Stakeholder",
            "res_model": "project.add.stakeholder.wizard",
            "res_id": self.id,
            "view_mode": "form",
            "context": dict(self.env.context),
        }
